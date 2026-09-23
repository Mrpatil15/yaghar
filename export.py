"""
export.py - Client delivery Excel spreadsheet generator and returned sheet re-importer.
"""

import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List, Tuple
from database import get_connection

def generate_client_excel(leads: List[Dict[str, Any]], batch_info: Dict[str, Any]) -> io.BytesIO:
    """
    Creates an executive, beautifully formatted Excel workbook for client delivery.
    """
    wb = openpyxl.Workbook()
    
    # ---------------- Sheet 1: Reactivated Leads ----------------
    ws = wb.active
    ws.title = "Calling Leads List"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="475569")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    
    fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_action_header = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    
    fill_hot = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")     # Emerald/Green light
    fill_warm = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")    # Amber light
    fill_cold = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")    # Sky light
    fill_status_col = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid") # Soft yellow editable

    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Title block
    client_name = batch_info.get("client_name", "Valued Client")
    batch_name = batch_info.get("batch_name", "Batch")
    corridor = batch_info.get("corridor", "Central Mumbai")
    
    ws["A1"] = f"Dead-Lead Reactivation Delivery — {client_name}"
    ws["A1"].font = font_title
    ws["A2"] = f"Batch: {batch_name} | Corridor: {corridor} | Priority calling list with customized Hinglish pitches"
    ws["A2"].font = font_subtitle

    # Table Headers
    headers = [
        "Lead Name",
        "Phone Number",
        "Score (0-100)",
        "Tier",
        "Assigned Hinglish Script",
        "Call Status (Team to Fill)",
        "Team Remarks / Next Visit",
        "Lead Source",
        "Original Enquiry Date"
    ]
    
    header_row_idx = 4
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=header_text)
        cell.font = font_header
        # Distinct color for columns the client's team needs to fill
        if "Team" in header_text:
            cell.fill = fill_action_header
        else:
            cell.fill = fill_header
        cell.alignment = align_center
        cell.border = cell_border
    
    ws.row_dimensions[header_row_idx].height = 28

    # Populate Rows
    row_idx = 5
    for lead in leads:
        tier_val = lead.get("tier", "Warm")
        score_val = lead.get("score")
        score_display = round(float(score_val), 1) if score_val is not None else ""
        
        # Format phone for display
        raw_phone = str(lead.get("phone", "")).strip()
        if len(raw_phone) == 10:
            formatted_phone = f"+91 {raw_phone[:5]} {raw_phone[5:]}"
        else:
            formatted_phone = raw_phone

        row_values = [
            lead.get("name", "Unknown"),
            formatted_phone,
            score_display,
            tier_val,
            lead.get("assigned_script", ""),
            lead.get("call_status", ""),
            lead.get("call_notes", ""),
            lead.get("source", ""),
            lead.get("enquiry_date", "")
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_body
            cell.border = cell_border

            # Tier specific styling
            if col_idx == 4:
                cell.alignment = align_center
                cell.font = font_bold
                if tier_val == "Hot":
                    cell.fill = fill_hot
                elif tier_val == "Warm":
                    cell.fill = fill_warm
                else:
                    cell.fill = fill_cold
            elif col_idx in (2, 3, 8, 9):
                cell.alignment = align_center
            elif col_idx == 5:
                cell.alignment = align_wrap
            elif col_idx in (6, 7):
                cell.alignment = align_center
                cell.fill = fill_status_col  # highlight editable action column
            else:
                cell.alignment = align_left

        ws.row_dimensions[row_idx].height = 42
        row_idx += 1

    # Column widths
    col_widths = {
        "A": 22,  # Name
        "B": 18,  # Phone
        "C": 14,  # Score
        "D": 12,  # Tier
        "E": 55,  # Assigned Script
        "F": 25,  # Call Status
        "G": 30,  # Remarks
        "H": 18,  # Source
        "I": 18   # Enquiry Date
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # ---------------- Sheet 2: Calling Guidelines ----------------
    ws_guide = wb.create_sheet(title="Calling Protocol & SOP")
    ws_guide.views.sheetView[0].showGridLines = True
    
    ws_guide["A1"] = "Real Estate Dead-Lead Calling Guidelines"
    ws_guide["A1"].font = font_title
    
    guidelines = [
        ("Tier Priority", "Call HOT tier leads on Day 1-2. Warm leads on Day 3-5. Cold leads via WhatsApp broadcast or Day 6-7."),
        ("Recommended Call Status Values", "Please fill Column F with: 'Connected', 'Converted' (Site Visit / Negotiation), 'No Answer', 'Not Interested', or 'Invalid'."),
        ("WhatsApp Follow-up", "If lead answers 'No Answer' after 2 attempts, dispatch the exact assigned Hinglish script via WhatsApp with your project brochure."),
        ("Re-importing Results", "Save this Excel file with updated Status and send back to your Reactivation Consultant to generate conversion metrics & ROI report.")
    ]
    
    for g_idx, (title, desc) in enumerate(guidelines, start=3):
        ws_guide.cell(row=g_idx, column=1, value=title).font = font_bold
        ws_guide.cell(row=g_idx, column=2, value=desc).font = font_body
        ws_guide.row_dimensions[g_idx].height = 24
        
    ws_guide.column_dimensions["A"].width = 30
    ws_guide.column_dimensions["B"].width = 80

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def reimport_client_status_sheet(batch_id: int, file_bytes: io.BytesIO) -> Dict[str, Any]:
    """
    Parses client returned Excel or CSV sheet and updates call_status and notes in database.
    """
    try:
        # Try reading as excel first, then csv
        try:
            df = pd.read_excel(file_bytes)
        except Exception:
            file_bytes.seek(0)
            df = pd.read_csv(file_bytes)
            
        # Normalize column names
        col_map = {str(col).strip().lower(): col for col in df.columns}
        
        # Look for phone and status columns
        phone_col = None
        status_col = None
        notes_col = None
        
        for k, original in col_map.items():
            if "phone" in k or "mobile" in k:
                phone_col = original
            elif "status" in k or "call status" in k:
                status_col = original
            elif "remark" in k or "note" in k or "next" in k:
                notes_col = original
                
        if not phone_col or not status_col:
            return {
                "success": False,
                "message": f"Could not find required 'Phone' or 'Status' column in uploaded sheet. Found columns: {list(df.columns)}"
            }
            
        conn = get_connection()
        updated_count = 0
        
        for _, row in df.iterrows():
            raw_phone = str(row.get(phone_col, "")).strip()
            # Clean phone digits
            digits = "".join(ch for ch in raw_phone if ch.isdigit())
            if len(digits) > 10:
                digits = digits[-10:] # take last 10 digits
                
            status_val = str(row.get(status_col, "")).strip()
            notes_val = str(row.get(notes_col, "")).strip() if notes_col else ""
            
            if digits and status_val and status_val.lower() != "nan":
                # Standardize status value
                std_status = "Not Called"
                s_lower = status_val.lower()
                if "convert" in s_lower or "visit" in s_lower or "booked" in s_lower:
                    std_status = "Converted"
                elif "connect" in s_lower or "talked" in s_lower or "interested" in s_lower and "not" not in s_lower:
                    std_status = "Connected"
                elif "no answer" in s_lower or "busy" in s_lower or "ringing" in s_lower or "not reachable" in s_lower:
                    std_status = "No Answer"
                elif "not interested" in s_lower or "drop" in s_lower or "cancel" in s_lower:
                    std_status = "Not Interested"
                else:
                    std_status = status_val.title()
                
                cur = conn.cursor()
                cur.execute("""
                    UPDATE leads
                    SET call_status = ?, call_notes = CASE WHEN ? != '' THEN ? ELSE call_notes END, last_updated = CURRENT_TIMESTAMP
                    WHERE batch_id = ? AND phone = ?
                """, (std_status, notes_val, notes_val, batch_id, digits))
                
                if cur.rowcount > 0:
                    updated_count += cur.rowcount
                    
        conn.commit()
        conn.close()
        
        return {
            "success": True,
            "updated_count": updated_count,
            "message": f"Successfully updated call status for {updated_count} leads!"
        }
    except Exception as e:
        return {
            "success": False,
            "message": f"Error parsing returned sheet: {str(e)}"
        }
